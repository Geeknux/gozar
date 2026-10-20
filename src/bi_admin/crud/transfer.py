"""CSV and XLSX export of any list, and import with a dry run.

Export runs the list's own query — row rules, filters, search, date hierarchy,
ordering and the chosen columns — and writes what the list shows: choice labels,
related rows by label (``Restricted`` for rows outside the user's scope), headers
in the user's language. ``CRUD_EXPORT_MAX_ROWS`` caps it (the response says when
it was cut off).

Import reads a CSV or XLSX file whose header row names fields (by name or label,
in English or Persian; the primary key column updates existing rows). Every row
is saved through ``save_object`` — the form's validation, relation scoping, the
out-of-scope check and the change log — inside one transaction. A dry run, or any
row with an error, rolls everything back; the report lists each row's outcome.
"""

from __future__ import annotations

import csv
import datetime as dt
import decimal
import io
import re
from typing import Any

from django.db import models, router, transaction
from django.utils import translation
from django.utils.encoding import force_str
from django.utils.text import capfirst
from django.utils.translation import gettext as _

from bi_admin.access.rules import model_queryset, target_queryset
from bi_admin.calendar import jalali
from bi_admin.conf import setting
from bi_admin.crud.service import (
    CrudError,
    apply_ordering,
    apply_select_related,
    chosen_columns,
    list_queryset,
    require,
    save_object,
    serialize_rows,
)
from bi_admin.discovery import model_label

FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
}
BATCH = 500
REPORT_ROWS = 500  # rows listed in an import report (totals count all)
TRUE = {"1", "true", "yes", "y", "on", "✔", "بله", "بلی", "درست", "x"}
FALSE = {"0", "false", "no", "n", "off", "✘", "خیر", "نه", "نادرست", ""}
PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


# --- export --------------------------------------------------------------------------


def _text(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, dict):
        return value.get("label") or value.get("name") or ""
    if isinstance(value, list):
        return ", ".join(str(_text(v)) for v in value)
    return value


def _excel_value(value: Any, kind: str) -> Any:
    value = _text(value)
    if value == "":
        return None
    if kind in ("integer", "decimal", "float") and isinstance(value, str):
        try:
            return decimal.Decimal(value) if kind == "decimal" else float(value)
        except (decimal.InvalidOperation, ValueError):
            return value
    if kind == "datetime" and isinstance(value, str):
        try:
            return dt.datetime.fromisoformat(value).replace(tzinfo=None)
        except ValueError:
            return value
    if kind == "date" and isinstance(value, str):
        try:
            return dt.date.fromisoformat(value)
        except ValueError:
            return value
    return value


def export_list(options, request, params: dict[str, str], fmt: str, pks=None):
    """``(content, content_type, filename, truncated)`` for ``fmt`` in ``FORMATS``."""
    require(options, request, "view")
    if not options.has_export_permission(request):
        raise CrudError(403, _("You may not export %(model)s") % {"model": options.label})
    if fmt not in FORMATS:
        raise CrudError(400, _("Unknown format %(format)r") % {"format": fmt})
    cols = chosen_columns(options, params)
    qs = list_queryset(options, request, params)
    if pks:
        qs = qs.filter(pk__in=pks)
    qs = apply_select_related(options, apply_ordering(options, qs, params.get("ordering")), cols)
    limit = setting("CRUD_EXPORT_MAX_ROWS") or None
    objs = list(qs[: limit + 1]) if limit else list(qs)
    truncated = bool(limit) and len(objs) > limit
    objs = objs[:limit] if limit else objs
    headers = [c.label for c in cols]
    rows = []
    for start in range(0, len(objs), BATCH):
        for row in serialize_rows(options, request, cols, objs[start : start + BATCH]):
            rows.append(row["cells"])
    kinds = [c.kind for c in cols]
    body = _csv(headers, rows) if fmt == "csv" else _xlsx(options, headers, rows, kinds)
    content_type, ext = FORMATS[fmt]
    filename = f"{options.info.model_name}.{ext}"
    return body, content_type, filename, truncated


def _csv(headers, rows) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_text(v) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: spreadsheet apps detect UTF-8


def _xlsx(options, headers, rows, kinds) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (
        re.sub(r"[\[\]:*?/\\]", "", capfirst(options.info.verbose_name_plural))[:31] or "data"
    )
    ws.sheet_view.rightToLeft = translation.get_language() == "fa"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_excel_value(v, k) for v, k in zip(row, kinds, strict=True)])
    for i, kind in enumerate(kinds, start=1):
        letter = get_column_letter(i)
        width = max([len(str(headers[i - 1]))] + [len(str(_text(r[i - 1]))) for r in rows[:200]])
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 60)
        fmt = {"datetime": "yyyy-mm-dd hh:mm", "date": "yyyy-mm-dd"}.get(kind)
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --- import --------------------------------------------------------------------------


def template(options, request, fmt: str):
    """An empty file with the importable columns (field names, so it round-trips)."""
    require(options, request, "view")
    if fmt not in FORMATS:
        raise CrudError(400, _("Unknown format %(format)r") % {"format": fmt})
    headers = [options.info.pk_name, *options.get_import_fields()]
    body = (
        _csv(headers, [])
        if fmt == "csv"
        else _xlsx(options, headers, [], ["string"] * len(headers))
    )
    content_type, ext = FORMATS[fmt]
    return body, content_type, f"{options.info.model_name}-import.{ext}"


def read_table(upload) -> list[list[Any]]:
    name = (getattr(upload, "name", "") or "").lower()
    raw = upload.read()
    if name.endswith(".xlsx") or raw[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise CrudError(400, _("The file is not a valid XLSX workbook")) from None
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("cp1256")  # Arabic/Persian Windows code page
        except UnicodeDecodeError:
            raise CrudError(400, _("The CSV file must be UTF-8 encoded")) from None
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def _norm(text: Any) -> str:
    return (
        re.sub(r"[\s_\-]+", " ", force_str(text or ""))
        .strip()
        .lower()
        .translate(str.maketrans({"ي": "ی", "ك": "ک"}))
    )


def _header_map(options) -> dict[str, str]:
    """Normalised header text → field name: names, verbose names in every language."""
    names = [options.info.pk_name, *options.get_import_fields()]
    out: dict[str, str] = {"pk": options.info.pk_name, "id": options.info.pk_name}
    for lang in ("en", "fa"):
        with translation.override(lang):
            for name in names:
                f = options.model._meta.get_field(name)
                out[_norm(f.verbose_name)] = name
    for name in names:
        out[_norm(name)] = name
    return out


def _parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = _norm(value)
    if text in TRUE:
        return True
    if text in FALSE:
        return False
    raise ValueError(_("Enter yes or no."))


def _parse_number(value: Any) -> Any:
    if isinstance(value, int | float | decimal.Decimal):
        return value
    text = force_str(value).translate(PERSIAN_DIGITS).replace("٬", "").replace(",", "").strip()
    return text.replace("٫", ".")


def _parse_date(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = force_str(value).strip().translate(PERSIAN_DIGITS)
    try:
        return jalali.parse_date(text[:10]).isoformat()
    except ValueError:
        return text


def _parse_datetime(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min).isoformat()
    text = force_str(value).strip().translate(PERSIAN_DIGITS)
    m = re.match(r"^(\d{4}[/-]\d{1,2}[/-]\d{1,2})(?:[ T](\d{1,2}:\d{2}(?::\d{2})?))?$", text)
    if m:
        try:
            day = jalali.parse_date(m.group(1))
        except ValueError:
            return text
        return f"{day.isoformat()} {m.group(2) or '00:00'}"
    return text


class _Resolver:
    """Relation cells: a primary key, or the label of exactly one row in scope."""

    def __init__(self, request):
        self.request = request
        self.cache: dict[tuple[str, str], Any] = {}

    def one(self, field, text: str):
        label = model_label(field.related_model)
        key = (label, text)
        if key in self.cache:
            return self.cache[key]
        qs = target_queryset(label, self.request)
        pk_field = field.related_model._meta.pk
        match = None
        try:
            pk = pk_field.to_python(text)
        except Exception:
            pk = None
        if pk is not None:
            match = qs.filter(pk=pk).values_list("pk", flat=True).first()
        if match is None:
            from bi_admin.crud.service import apply_search, get_registry, site

            candidates = qs
            if label in get_registry():
                candidates = apply_search(site.options_for(label), qs, text)
            hits = [o.pk for o in candidates[:50] if _norm(force_str(o)) == _norm(text)]
            if len(hits) > 1:
                raise ValueError(_("“%(value)s” matches more than one row.") % {"value": text})
            match = hits[0] if hits else None
        if match is None:
            raise ValueError(_("“%(value)s” was not found.") % {"value": text})
        self.cache[key] = match
        return match


def _choice(field, value: Any):
    text = _norm(value)
    for lang in ("en", "fa"):
        with translation.override(lang):
            for v, label in field.flatchoices:
                if _norm(v) == text or _norm(label) == text:
                    return v
    return value


def _cell(field, value: Any, resolver: _Resolver) -> Any:
    if isinstance(value, str):
        value = value.strip()
    if value is None or value == "":
        if isinstance(field, models.BooleanField):
            return False
        if field.many_to_many:
            return []
        if field.null or isinstance(field, models.DateField | models.DecimalField):
            return None
        if field.is_relation or isinstance(field, models.IntegerField | models.FloatField):
            return None
        return ""
    if field.many_to_many:
        parts = [p.strip() for p in re.split(r"[;\n|]|,(?=\s*\S)", force_str(value)) if p.strip()]
        return [resolver.one(field, p) for p in parts]
    if field.is_relation:
        return resolver.one(field, force_str(value))
    if field.choices:
        return _choice(field, value)
    if isinstance(field, models.BooleanField):
        return _parse_bool(value)
    if isinstance(field, models.DateTimeField):
        return _parse_datetime(value)
    if isinstance(field, models.DateField):
        return _parse_date(value)
    if isinstance(field, models.IntegerField | models.DecimalField | models.FloatField):
        return _parse_number(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return force_str(value)


def import_file(options, request, upload, dry_run: bool = True) -> dict[str, Any]:
    if not options.has_import_permission(request):
        raise CrudError(403, _("You may not import %(model)s") % {"model": options.label})
    if upload is None:
        raise CrudError(400, _("Choose a CSV or XLSX file"))
    table = [r for r in read_table(upload) if any(c not in (None, "") for c in r)]
    if not table:
        raise CrudError(400, _("The file is empty"))
    header, body = table[0], table[1:]
    limit = setting("CRUD_IMPORT_MAX_ROWS")
    if limit and len(body) > limit:
        raise CrudError(
            400,
            _("The file has %(rows)s rows; the limit is %(limit)s.")
            % {"rows": len(body), "limit": limit},
            code="too_many_rows",
        )
    mapping = _header_map(options)
    columns, ignored = [], []
    seen = set()
    for i, text in enumerate(header):
        name = mapping.get(_norm(text))
        if name is None or name in seen:
            if text not in (None, ""):
                ignored.append(force_str(text))
            continue
        seen.add(name)
        columns.append((i, name))
    pk_name = options.info.pk_name
    fields = [n for _i, n in columns if n != pk_name]
    if not fields:
        raise CrudError(400, _("No column of the file matches an importable field"))
    resolver = _Resolver(request)
    can_add = options.has_add_permission(request)
    can_change = options.has_change_permission(request)
    scoped = model_queryset(options, request)
    pk_field = options.model._meta.pk
    results: list[dict[str, Any]] = []
    totals = {"create": 0, "update": 0, "skip": 0, "error": 0}

    with transaction.atomic(using=router.db_for_write(options.model)):
        for offset, raw in enumerate(body, start=2):
            values: dict[str, Any] = {}
            errors: dict[str, list[str]] = {}
            pk_text = None
            for i, name in columns:
                cell = raw[i] if i < len(raw) else None
                if name == pk_name:
                    pk_text = None if cell in (None, "") else force_str(cell).strip()
                    if isinstance(cell, float) and cell.is_integer():
                        pk_text = str(int(cell))
                    continue
                try:
                    values[name] = _cell(options.model._meta.get_field(name), cell, resolver)
                except ValueError as exc:
                    errors.setdefault(name, []).append(force_str(exc.args[0] if exc.args else exc))
            obj = None
            if pk_text:
                try:
                    obj = scoped.filter(pk=pk_field.to_python(pk_text)).first()
                except Exception:
                    obj = None
                if obj is None:
                    errors.setdefault("__all__", []).append(
                        _("No record with ID %(pk)s that you can change.") % {"pk": pk_text}
                    )
            action = "update" if obj is not None else "create"
            if not errors and action == "create" and not can_add:
                errors["__all__"] = [_("You may not add records.")]
            if not errors and action == "update" and not can_change:
                errors["__all__"] = [_("You may not change records.")]
            entry: dict[str, Any] = {"row": offset, "action": action}
            if not errors:
                try:
                    with transaction.atomic(using=router.db_for_write(options.model)):
                        before = obj.pk if obj is not None else None
                        saved = save_object(
                            options,
                            request,
                            values,
                            obj,
                            partial=True,
                            fields=fields,
                            log_extra={"source": "import"},
                        )
                        if getattr(saved, "_gozar_unchanged", False):
                            entry["action"] = "skip"  # nothing changed
                        entry["pk"] = None if dry_run and before is None else force_str(saved.pk)
                        entry["label"] = force_str(saved)
                except CrudError as exc:
                    if exc.status not in (400, 403):
                        raise
                    errors = exc.payload.get("errors") or {"__all__": [exc.payload["detail"]]}
            if errors:
                entry["action"] = "error"
                entry["errors"] = errors
            totals[entry["action"]] += 1
            if len(results) < REPORT_ROWS:
                results.append(entry)
        committed = not dry_run and not totals["error"]
        if not committed:
            transaction.set_rollback(True)
    labels = {}
    for _i, name in columns:
        labels[name] = capfirst(force_str(options.model._meta.get_field(name).verbose_name))
    return {
        "dry_run": dry_run,
        "committed": committed,
        "totals": totals,
        "rows": results,
        "row_count": len(body),
        "columns": [{"field": n, "label": labels[n]} for _i, n in columns],
        "ignored_columns": ignored,
    }
