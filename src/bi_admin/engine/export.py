"""Export a report result to CSV or XLSX, with headers in the active language."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from django.utils.text import slugify

from bi_admin.engine.runner import ReportResult
from bi_admin.semantic.base import active_language

FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
}
XLSX_NUMBER_FORMATS = {
    "currency": "#,##0",
    "integer": "#,##0",
    "percent": "0.00%",
    "decimal": "#,##0.00",
    "number": "#,##0.##",
}


def _data(result: ReportResult | dict) -> dict[str, Any]:
    """A run result or its stored ``as_dict()`` form (background jobs)."""
    return result.as_dict() if isinstance(result, ReportResult) else result


def table(result: ReportResult | dict) -> tuple[list[str], list[list[Any]], list[dict]]:
    """Header labels, row values and column metadata for an export.

    Time buckets export their label (``2026-10``, ``1405-S3``…) and choice
    values their display label, both as shown in the UI."""
    data = _data(result)
    columns = data["columns"]
    headers = [c["label"] for c in columns]
    rows = []
    for row in data["rows"]:
        values = []
        for c in columns:
            value = row.get(c["key"])
            if c.get("grain"):
                value = row.get(c["label_key"])
            elif c.get("choices") and value is not None:
                value = next((ch["label"] for ch in c["choices"] if ch["value"] == value), value)
            values.append(value)
        rows.append(values)
    return headers, rows, columns


def to_csv(result: ReportResult | dict) -> bytes:
    headers, rows, _ = table(result)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(["" if v is None else v for v in row] for row in rows)
    # BOM so spreadsheet apps detect UTF-8 (Persian headers and values)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def sheet_title(title: str) -> str:
    """``title`` as an Excel sheet name (31 characters, no ``[]:*?/\\``)."""
    return re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:31].strip()


def to_xlsx(result: ReportResult | dict, title: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    headers, rows, columns = table(result)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title(title) or _data(result)["semantic_model"][:31]
    ws.sheet_view.rightToLeft = active_language() == "fa"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, col in enumerate(columns, start=1):
        letter = get_column_letter(i)
        width = max([len(str(headers[i - 1]))] + [len(str(r[i - 1] or "")) for r in rows[:200]])
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 60)
        if col["role"] in ("measure", "comparison"):
            fmt = XLSX_NUMBER_FORMATS.get(col.get("format", "number"), "General")
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export(result: ReportResult | dict, fmt: str, title: str = "") -> tuple[bytes, str, str]:
    """``(content, content_type, filename)`` for ``fmt`` in ``FORMATS``. ``title``
    (a saved report's or widget's name, in the viewer's language) names the file
    and the sheet; without it, the semantic model does."""
    data = _data(result)
    content_type, ext = FORMATS[fmt]
    body = to_csv(data) if fmt == "csv" else to_xlsx(data, title)
    stem = slugify(title, allow_unicode=True)[:100] or data["semantic_model"]
    filename = f"{stem}.{ext}"
    return body, content_type, filename
