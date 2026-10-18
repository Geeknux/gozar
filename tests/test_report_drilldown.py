"""M3: drill-down, export and dimension values for the report UI.

Numbers come from the hand-made ``orders`` fixture in conftest.py.
"""

import csv
import io

import pytest
from django.db.models import Count
from django.test import Client
from openpyxl import load_workbook

from bi_admin.engine import run
from bi_admin.semantic import Dimension, Measure, Semantic, register, semantic_site
from shop.models import Order

API = "/gozar/api/reports"
SPEC = {"semantic_model": "order_metrics", "measures": ["orders", "revenue"]}

pytestmark = pytest.mark.django_db


def post(client, path, body, **params):
    query = "&".join(f"{k}={v}" for k, v in params.items())
    return client.post(f"{API}/{path}{'?' + query if query else ''}", body, "application/json")


def drill(client, cell, spec=None, **extra):
    return post(client, "drilldown", {"spec": {**SPEC, **(spec or {})}, "cell": cell, **extra})


def amounts(body):
    """Order amounts on a drill-down page (the "amount" list column)."""
    i = [c["name"] for c in body["columns"]].index("amount")
    return sorted(row["cells"][i] for row in body["rows"])


# --- drill-down: which records ---------------------------------------------------------


def test_drill_into_a_category(admin_client, orders):
    body = drill(admin_client, {"city": "تهران"}, {"dimensions": ["city"]}).json()
    assert body["count"] == 5 and body["model"] == "shop.order" and body["linkable"]
    assert amounts(body) == ["100", "1000", "200", "50", "500"]
    # clicking the revenue cell applies revenue's own filter (paid orders only)
    paid = drill(admin_client, {"city": "تهران"}, {"dimensions": ["city"]}, measure="revenue")
    assert paid.json()["count"] == 4


def test_counts_match_the_report_cells(admin_client, orders):
    spec = {**SPEC, "measures": ["orders", "paid_orders"], "dimensions": ["province", "status"]}
    report = run(spec).rows
    assert len(report) == 6
    for row in report:
        cell = {"province": row["province"], "status": row["status"]}
        assert drill(admin_client, cell, spec).json()["count"] == row["orders"]
        paid = drill(admin_client, cell, spec, measure="paid_orders").json()
        assert paid["count"] == row["paid_orders"]


@pytest.mark.parametrize(
    "calendar, grain, start, expected",
    [
        ("gregorian", "month", "2026-09-01", ["100", "200", "50"]),  # September
        ("gregorian", "season", "2026-07-01", ["100", "200", "400", "50", "70"]),  # Q3
        ("gregorian", "week", "2026-09-21", ["100", "200"]),  # ISO week from Monday
        ("jalali", "month", "2026-09-23", ["200", "300", "50"]),  # Mehr 1405
        ("jalali", "month", "1405/06/01", ["100", "70"]),  # Shahrivar, Jalali notation
        ("jalali", "season", "2026-06-22", ["100", "400", "70"]),  # summer 1405
        ("jalali", "year", "2025-03-21", ["1000", "500"]),  # 1404
        ("jalali", "week", "2026-09-19", ["100", "200"]),  # Saturday-start week
        ("jalali", "day", "2026-09-23", ["200"]),  # local date (00:30 Tehran)
    ],
)
def test_drill_into_time_buckets(admin_client, orders, calendar, grain, start, expected):
    spec = {"time_grain": grain, "calendar": calendar}
    body = drill(admin_client, {"created_at": start}, spec).json()
    assert amounts(body) == expected


def test_time_buckets_match_the_report(admin_client, orders, calendar_table):
    spec = {**SPEC, "time_grain": "month", "calendar": "jalali", "dimensions": ["city"]}
    for row in run(spec).rows:
        cell = {"created_at": row["created_at"], "city": row["city"]}
        assert drill(admin_client, cell, spec).json()["count"] == row["orders"]


def test_null_bucket_and_report_filters(admin_client, orders):
    spec = {"time_dimension": "paid_at", "time_grain": "year"}
    assert drill(admin_client, {"paid_at": None}, spec).json()["count"] == 2  # never paid
    filtered = {
        "dimensions": ["city"],
        "filters": [{"field": "status", "op": "eq", "value": "paid"}],
    }
    assert drill(admin_client, {"city": "تهران"}, filtered).json()["count"] == 4
    assert drill(admin_client, {}, filtered).json()["count"] == 6  # the whole report


def test_relative_period_is_kept(admin_client, orders, frozen_today):
    spec = {"dimensions": ["city"], "relative_period": "last_month", "calendar": "jalali"}
    body = drill(admin_client, {"city": "تهران"}, spec).json()
    assert amounts(body) == ["100"]  # Shahrivar 1405, Tehran: order #1 only


def test_choice_labels_are_accepted(admin_client, orders):
    spec = {"dimensions": ["status"]}
    assert drill(admin_client, {"status": "Refunded"}, spec).json()["count"] == 1


# --- drill-down: errors, permissions, pagination -------------------------------------


@pytest.mark.parametrize(
    "cell, spec, extra, code, path",
    [
        ({"province": "Fars"}, {"dimensions": ["city"]}, {}, "not_a_group_key", "cell.province"),
        ({"status": "lost"}, {"dimensions": ["status"]}, {}, "invalid_choice", "cell.status"),
        (
            {"created_at": "2026-09-15"},
            {"time_grain": "month"},
            {},
            "invalid_value",
            "cell.created_at",
        ),
        ({"created_at": "soon"}, {"time_grain": "month"}, {}, "invalid_value", "cell.created_at"),
        ({}, {}, {"measure": "aov"}, "unknown_measure", "measure"),
        ({}, {"measures": ["nope"]}, {}, "unknown_measure", "measures[0]"),
    ],
)
def test_drilldown_errors(admin_client, orders, cell, spec, extra, code, path):
    resp = drill(admin_client, cell, spec, **extra)
    assert resp.status_code == 422
    assert [(e["code"], e["path"]) for e in resp.json()["errors"]] == [(code, path)]


def test_drilldown_permissions(client, make_staff, orders):
    assert drill(client, {}).status_code == 401
    client.force_login(make_staff("shop.view_customer", username="nope"))
    resp = drill(client, {})
    assert resp.status_code == 403 and resp.json()["code"] == "permission_denied"
    client.force_login(make_staff("shop.view_order"))
    assert drill(client, {}).json()["count"] == 9


def test_drilldown_respects_the_semantic_queryset(admin_client, orders, monkeypatch):
    sem = semantic_site.get("order_metrics")
    monkeypatch.setattr(sem, "get_queryset", lambda request: Order.objects.filter(amount__gte=300))
    assert drill(admin_client, {}).json()["count"] == 5


def test_drilldown_pagination_and_rows(admin_client, orders):
    first = drill(admin_client, {}, page_size=4).json()
    assert (first["count"], first["num_pages"], len(first["rows"])) == (9, 3, 4)
    last = drill(admin_client, {}, page_size=4, page=3).json()
    assert len(last["rows"]) == 1
    pks = {r["pk"] for r in first["rows"]} | {r["pk"] for r in last["rows"]}
    assert pks <= set(Order.objects.values_list("pk", flat=True))
    assert first["rows"][0]["label"].startswith("#")
    names = [c["name"] for c in first["columns"]]
    assert names == [
        "__str__",
        "customer",
        "customer__city__name",
        "status",
        "amount",
        "created_at",
    ]
    assert drill(admin_client, {}, page_size=4, page=9).status_code == 404


def test_drilldown_on_a_relation_dimension(admin_client, orders):
    class ByCustomer(Semantic):
        model = Order
        dimensions = {"customer": Dimension("customer")}
        measures = {"n": Measure(Count("id"))}

    register(ByCustomer)
    try:
        spec = {"semantic_model": "by_customer", "measures": ["n"], "dimensions": ["customer"]}
        rows = run(spec).rows
        pk = orders[0].customer_id
        assert next(r for r in rows if r["customer"] == pk)["n"] == 3
        body = post(admin_client, "drilldown", {"spec": spec, "cell": {"customer": pk}}).json()
        assert body["count"] == 3
        values = admin_client.get(
            f"{API}/values", {"semantic_model": "by_customer", "dimension": "customer", "q": "علی"}
        ).json()["values"]
        assert values == [{"value": pk, "label": "علی"}]
    finally:
        semantic_site.unregister("by_customer")


# --- export -----------------------------------------------------------------------


def read_csv(resp):
    text = resp.content.decode("utf-8")
    assert text.startswith("﻿")
    return list(csv.reader(io.StringIO(text[1:])))


def test_csv_export_matches_the_report(admin_client, orders):
    spec = {**SPEC, "dimensions": ["city"]}
    resp = post(admin_client, "export", spec, format="csv")
    assert resp.status_code == 200
    assert resp["Content-Type"].startswith("text/csv")
    assert resp["Content-Disposition"] == 'attachment; filename="order_metrics.csv"'
    rows = read_csv(resp)
    assert rows[0] == ["City", "Orders", "Revenue"]
    assert rows[1:] == [["تهران", "5", "1800"], ["شیراز", "2", "300"], ["مشهد", "2", "600"]]
    report = run(spec).rows
    assert sum(int(r[1]) for r in rows[1:]) == sum(r["orders"] for r in report) == 9


def test_export_time_buckets_and_choice_labels(admin_client, orders):
    spec = {
        **SPEC,
        "dimensions": ["status"],
        "time_grain": "year",
        "sort": [{"field": "created_at", "direction": "asc"}],
    }
    rows = read_csv(post(admin_client, "export", spec, format="csv"))
    assert rows[0] == ["Order date", "Status", "Orders", "Revenue"]
    assert ["2025", "Paid", "1", "1000"] in rows
    assert ["2026", "Awaiting payment", "1", ""] in rows  # no paid rows → empty revenue


def test_export_headers_follow_the_users_language(admin_client, orders):
    admin_client.patch("/gozar/api/preferences", {"language": "fa"}, "application/json")
    spec = {**SPEC, "dimensions": ["status"]}
    rows = read_csv(post(admin_client, "export", spec, format="csv"))
    assert rows[0] == ["وضعیت", "تعداد سفارش", "درآمد"]
    assert ["پرداخت‌شده", "6", "2700"] in rows

    resp = post(admin_client, "export", spec, format="xlsx")
    assert resp["Content-Disposition"].endswith('.xlsx"')
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws.sheet_view.rightToLeft is True
    assert [c.value for c in ws[1]] == ["وضعیت", "تعداد سفارش", "درآمد"]
    data = {r[0]: r[1:] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert data["پرداخت‌شده"] == (6, 2700)
    assert ws["B2"].number_format == "#,##0" and ws["C2"].number_format == "#,##0"


def test_xlsx_export_in_english(admin_client, orders):
    spec = {**SPEC, "measures": ["aov"], "dimensions": ["city"]}
    ws = load_workbook(io.BytesIO(post(admin_client, "export", spec, format="xlsx").content)).active
    assert ws.sheet_view.rightToLeft is not True
    assert [c.value for c in ws[1]] == ["City", "Avg order value"]
    assert {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)} == {
        "تهران": 450,  # (100 + 200 + 500 + 1000) / 4
        "شیراز": 300,
        "مشهد": 600,
    }


def test_export_permissions_and_validation(client, make_staff, orders):
    assert post(client, "export", SPEC, format="csv").status_code == 401
    client.force_login(make_staff("shop.view_customer", username="nope"))
    assert post(client, "export", SPEC, format="csv").status_code == 403
    client.force_login(make_staff("shop.view_order"))
    assert post(client, "export", SPEC, format="csv").status_code == 200
    assert post(client, "export", SPEC, format="pdf").status_code == 422
    bad = post(client, "export", {**SPEC, "measures": ["nope"]}, format="xlsx")
    assert bad.status_code == 422 and bad.json()["errors"][0]["code"] == "unknown_measure"


def test_export_requires_csrf(admin_user, orders):
    c = Client(enforce_csrf_checks=True)
    c.force_login(admin_user)
    assert post(c, "export", SPEC, format="csv").status_code in (401, 403)


# --- dimension values ---------------------------------------------------------------


def values(client, dimension, q="", model="order_metrics"):
    return client.get(f"{API}/values", {"semantic_model": model, "dimension": dimension, "q": q})


def test_values_for_choices_and_text(admin_client, orders):
    status = values(admin_client, "status").json()["values"]
    assert [v["value"] for v in status] == ["pending", "paid", "cancelled", "refunded"]
    assert status[1] == {"value": "paid", "label": "Paid"}
    assert [v["value"] for v in values(admin_client, "status", "pai").json()["values"]] == ["paid"]
    cities = values(admin_client, "city").json()
    assert [v["value"] for v in cities["values"]] == ["تهران", "شیراز", "مشهد"]
    assert cities["more"] is False
    # Arabic yeh in the query still finds Persian text
    assert [v["value"] for v in values(admin_client, "city", "شير").json()["values"]] == ["شیراز"]
    amounts_ = [v["value"] for v in values(admin_client, "amount").json()["values"]]
    assert amounts_ == ["50", "70", "100", "200", "300", "400", "500", "600", "1000"]


def test_values_follow_the_language_and_permissions(client, make_staff, orders):
    assert values(client, "status").status_code == 401
    client.force_login(make_staff("shop.view_customer", username="nope"))
    assert values(client, "status").status_code == 403
    client.force_login(make_staff("shop.view_order"))
    client.patch("/gozar/api/preferences", {"language": "fa"}, "application/json")
    assert values(client, "status", "مرجوع").json()["values"] == [
        {"value": "refunded", "label": "مرجوعی"}
    ]
    assert values(client, "created_at").status_code == 422  # time dimensions use date pickers
    assert values(client, "nope").status_code == 422
