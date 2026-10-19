"""M6: scheduled e-mail reports — due times, per-recipient content, delivery log."""

import csv
import datetime as dt
import io
from io import StringIO

import pytest
from conftest import API, SHOP_VIEW, TEHRAN_REGION
from django.core import mail
from django.core.management import call_command
from django.test import Client
from django.utils import timezone

from bi_admin import scheduling
from bi_admin.models import (
    Dashboard,
    DashboardWidget,
    ReportLog,
    SavedReport,
    Schedule,
    ScheduleDelivery,
    UserPreferences,
)

UTC = dt.UTC
SCHEDULES = f"{API}/schedules"
BY_CITY = {
    "semantic_model": "order_metrics",
    "measures": ["revenue", "orders"],
    "dimensions": ["city"],
    "chart": "bar",
}
EAST_REGION = ["Mashhad", "Kerman"]
ALL_CITIES = [
    *("Tehran", "Mashhad", "Isfahan", "Shiraz", "Tabriz"),
    *("Karaj", "Ahvaz", "Qom", "Rasht", "Kerman"),
]


# --- due times -----------------------------------------------------------------------------


def at(*args, tz="UTC"):
    from zoneinfo import ZoneInfo

    return dt.datetime(*args, tzinfo=ZoneInfo(tz))


@pytest.mark.parametrize(
    ("after", "expected"),
    [
        # 08:00 Tehran = 04:30 UTC (no DST since 2022)
        (at(2026, 10, 19, 4, 0), at(2026, 10, 19, 4, 30)),
        (at(2026, 10, 19, 4, 30), at(2026, 10, 20, 4, 30)),  # strictly after
        (at(2026, 10, 19, 23, 0), at(2026, 10, 20, 4, 30)),
        # 21:00 UTC on the 19th is already the 20th in Tehran (00:30)
        (at(2026, 10, 19, 21, 0), at(2026, 10, 20, 4, 30)),
    ],
)
def test_daily(after, expected):
    nxt = scheduling.compute_next("daily", dt.time(8, 0), "Asia/Tehran", after)
    assert nxt == expected and nxt.tzinfo is UTC


def test_weekly():
    # 2026-10-19 is a Monday; Saturday = 5
    after = at(2026, 10, 19, 12, 0)
    nxt = scheduling.compute_next("weekly", dt.time(9, 0), "UTC", after, weekday=5)
    assert nxt == at(2026, 10, 24, 9, 0)
    same_day = scheduling.compute_next("weekly", dt.time(13, 0), "UTC", after, weekday=0)
    assert same_day == at(2026, 10, 19, 13, 0)
    next_week = scheduling.compute_next("weekly", dt.time(11, 0), "UTC", after, weekday=0)
    assert next_week == at(2026, 10, 26, 11, 0)


@pytest.mark.parametrize(
    ("after", "day", "expected"),
    [
        (at(2026, 1, 31, 10, 0), 31, at(2026, 2, 28, 9, 0)),  # clamped to February
        (at(2028, 2, 1, 0, 0), 31, at(2028, 2, 29, 9, 0)),  # leap year
        (at(2026, 2, 28, 9, 0), 31, at(2026, 3, 31, 9, 0)),
        (at(2026, 10, 19, 0, 0), 1, at(2026, 11, 1, 9, 0)),
        (at(2026, 12, 15, 0, 0), 15, at(2026, 12, 15, 9, 0)),
        (at(2026, 12, 31, 10, 0), 1, at(2027, 1, 1, 9, 0)),  # across the year boundary
    ],
)
def test_monthly_gregorian(after, day, expected):
    assert scheduling.compute_next("monthly", dt.time(9, 0), "UTC", after, day_of_month=day) == (
        expected
    )


@pytest.mark.parametrize(
    ("after", "day", "expected"),
    [
        # 1405/07/27 → the 1st of Aban 1405 = 2026-10-23
        (at(2026, 10, 19), 1, at(2026, 10, 23, 9, 0)),
        # day 31 in Mehr (30 days) is its last day, 1405/07/30 = 2026-10-22
        (at(2026, 10, 19), 31, at(2026, 10, 22, 9, 0)),
        # Esfand 1403 (leap) has a 30th: 2025-03-20; Esfand 1404 ends on the 29th
        (at(2025, 3, 1), 30, at(2025, 3, 20, 9, 0)),
        (at(2026, 3, 1), 30, at(2026, 3, 20, 9, 0)),
        # across Nowruz: after 1404/12/29 comes 1405/01/01 = 2026-03-21
        (at(2026, 3, 20, 10, 0), 1, at(2026, 3, 21, 9, 0)),
    ],
)
def test_monthly_jalali(after, day, expected):
    nxt = scheduling.compute_next(
        "monthly", dt.time(9, 0), "UTC", after, day_of_month=day, calendar="jalali"
    )
    assert nxt == expected


def test_daylight_saving_time():
    # New York: 08:00 is 12:00 UTC in summer and 13:00 UTC in winter
    summer = scheduling.compute_next("daily", dt.time(8), "America/New_York", at(2026, 10, 30))
    winter = scheduling.compute_next("daily", dt.time(8), "America/New_York", at(2026, 11, 2))
    assert (summer, winter) == (at(2026, 10, 30, 12), at(2026, 11, 2, 13))


def test_unknown_timezone_falls_back_to_the_project_one():
    assert scheduling.zone("Mars/Olympus").key == "Asia/Tehran"
    assert not scheduling.is_timezone("Mars/Olympus") and scheduling.is_timezone("Europe/Berlin")


# --- fixtures ----------------------------------------------------------------------------


@pytest.fixture
def people(seeded, make_staff, regional):
    users = {
        "owner": make_staff(*SHOP_VIEW, username="owner"),
        "tehran": regional("Sales: Tehran region"),
        "east": regional("Sales: East region", username="east"),
        "outsider": make_staff(username="outsider"),
    }
    for name, user in users.items():
        user.email = f"{name}@example.com"
        user.save()
    return users


@pytest.fixture
def report(people):
    return SavedReport.objects.create(
        owner=people["owner"],
        name="Revenue by city",
        description="All cities",
        semantic_model="order_metrics",
        spec=BY_CITY,
        visibility="everyone",
    )


@pytest.fixture
def schedule(people, report):
    s = Schedule.objects.create(
        owner=people["owner"],
        saved_report=report,
        frequency="daily",
        time=dt.time(8, 0),
        timezone="Asia/Tehran",
        attachment="csv",
    )
    s.recipients.set([people["owner"], people["tehran"], people["east"]])
    s.next_run_at = timezone.now() - dt.timedelta(minutes=1)
    s.save()
    return s


def message_for(username):
    (msg,) = [m for m in mail.outbox if m.to == [f"{username}@example.com"]]
    return msg


def html(msg) -> str:
    (alternative,) = msg.alternatives
    return alternative[0] if isinstance(alternative, tuple) else alternative.content


def csv_cities(msg) -> set[str]:
    (attachment,) = [a for a in msg.attachments if a[0] == "revenue-by-city.csv"]
    name, content, mimetype = attachment
    assert mimetype == "text/csv"
    text = content.decode("utf-8-sig") if isinstance(content, bytes) else content.lstrip("﻿")
    rows = list(csv.reader(io.StringIO(text)))
    return {r[0] for r in rows[1:]}


# --- sending ------------------------------------------------------------------------------


def test_each_recipient_gets_their_own_scope(schedule):
    deliveries = scheduling.send_due()
    assert [d.status for d in deliveries] == ["sent"] * 3
    assert len(mail.outbox) == 3
    tehran, east, owner = message_for("tehran"), message_for("east"), message_for("owner")
    tehran_cities, east_cities = csv_cities(tehran), csv_cities(east)
    assert tehran_cities and tehran_cities <= set(TEHRAN_REGION)
    assert east_cities and east_cities <= set(EAST_REGION)
    assert csv_cities(owner) > tehran_cities | east_cities
    for city in set(ALL_CITIES) - set(TEHRAN_REGION):
        assert f">{city}<" not in html(tehran) and city not in tehran.body
    for city in tehran_cities:
        assert f">{city}<" in html(tehran)
    assert tehran.subject.startswith("Revenue by city · ")
    assert "Open in Gozar" in html(tehran)


def test_due_schedules_are_sent_once_and_move_on(schedule):
    before = schedule.next_run_at
    scheduling.send_due()
    assert len(mail.outbox) == 3
    schedule.refresh_from_db()
    assert schedule.next_run_at > timezone.now() and schedule.last_run_at is not None
    assert scheduling.send_due() == []  # not due any more
    assert len(mail.outbox) == 3
    # a process that read the old next_run_at can't claim it again
    stale = Schedule.objects.get(pk=schedule.pk)
    stale.next_run_at = before
    assert scheduling.claim(stale, timezone.now()) is False


def test_disabled_and_future_schedules_wait(schedule):
    Schedule.objects.filter(pk=schedule.pk).update(enabled=False)
    assert scheduling.send_due() == []
    Schedule.objects.filter(pk=schedule.pk).update(
        enabled=True, next_run_at=timezone.now() + dt.timedelta(hours=1)
    )
    assert scheduling.send_due() == []
    assert len(scheduling.send_due(timezone.now() + dt.timedelta(hours=2))) == 3


def test_language_calendar_and_xlsx_attachment(schedule, people, calendar_table):
    from openpyxl import load_workbook

    UserPreferences.objects.create(user=people["tehran"], language="fa", calendar="jalali")
    schedule.attachment = "xlsx"
    schedule.save()
    SavedReport.objects.filter(pk=schedule.saved_report_id).update(
        spec={
            **BY_CITY,
            "dimensions": [],
            "time_grain": "month",
            "relative_period": "past_12_months",
        }
    )
    scheduling.send_due(timezone.now())
    fa, en = message_for("tehran"), message_for("east")
    body = html(fa)
    assert 'lang="fa" dir="rtl"' in body
    assert "تعداد سفارش" in body and "درآمد" in body  # Persian column labels
    assert ">Revenue<" not in body
    assert "باز کردن در گذر" in body  # "Open in Gozar"
    assert any("۰" <= ch <= "۹" for ch in fa.subject)  # Persian digits
    # months are labelled in the Jalali calendar
    assert "مهر ۱۴۰۵" in body
    assert 'lang="en" dir="ltr"' in html(en) and "Revenue" in html(en)
    (name, content, mimetype) = fa.attachments[0]
    assert name.endswith(".xlsx") and "spreadsheetml" in mimetype
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet.sheet_view.rightToLeft is True
    assert [c.value for c in sheet[1]][:2] == ["تاریخ ثبت سفارش", "درآمد"]
    assert str(sheet["A2"].value).startswith("140")  # a Jalali bucket label
    en_sheet = load_workbook(io.BytesIO(en.attachments[0][1])).active
    assert [c.value for c in en_sheet[1]][:2] == ["Order date", "Revenue"]


def test_kpi_comparison_in_the_email(schedule):
    SavedReport.objects.filter(pk=schedule.saved_report_id).update(
        spec={
            "semantic_model": "order_metrics",
            "measures": ["orders"],
            "relative_period": "past_12_months",
            "compare_to": "same_period_last_year",
            "chart": "kpi",
        }
    )
    scheduling.send_due()
    body = html(message_for("owner"))
    assert "kpi-value" in body and "was" in body
    assert "compared with the same period last year" in body
    assert any(arrow in body for arrow in ("▲", "▼", "■"))


def test_every_send_is_audited(schedule, people):
    scheduling.send_due()
    deliveries = ScheduleDelivery.objects.order_by("username")
    assert [(d.username, d.status, d.reports, d.attachments) for d in deliveries] == [
        ("east", "sent", 1, 1),
        ("owner", "sent", 1, 1),
        ("tehran", "sent", 1, 1),
    ]
    logs = ReportLog.objects.filter(kind="email").order_by("username")
    assert [(log.username, log.status) for log in logs] == [
        ("east", "ok"),
        ("owner", "ok"),
        ("tehran", "ok"),
    ]
    assert "Tehran" in logs.get(username="tehran").scope or logs.get(username="tehran").scope
    assert logs.get(username="owner").scope == ""  # unrestricted


def test_recipients_without_email_or_access_are_skipped(schedule, people):
    people["east"].email = ""
    people["east"].save()
    SavedReport.objects.filter(pk=schedule.saved_report_id).update(visibility="private")
    deliveries = {d.username: d for d in scheduling.send_due()}
    assert (deliveries["east"].status, deliveries["east"].error) == ("skipped", "no_email_address")
    assert (deliveries["tehran"].status, deliveries["tehran"].error) == ("skipped", "not_shared")
    assert deliveries["owner"].status == "sent"
    assert [m.to for m in mail.outbox] == [["owner@example.com"]]


def test_a_failing_recipient_does_not_stop_the_others(schedule, monkeypatch):
    from django.core.mail import EmailMultiAlternatives

    original = EmailMultiAlternatives.send

    def send(self, *args, **kwargs):
        if self.to == ["tehran@example.com"]:
            raise ConnectionRefusedError("SMTP down")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(EmailMultiAlternatives, "send", send)
    deliveries = {d.username: d for d in scheduling.send_due()}
    assert deliveries["tehran"].status == "failed" and "SMTP down" in deliveries["tehran"].error
    assert deliveries["east"].status == deliveries["owner"].status == "sent"


def test_dashboard_email_has_a_section_per_widget(people, report):
    board = Dashboard.objects.create(
        owner=people["owner"], name="Sales overview", visibility="everyone",
        filters={"values": {"city": ["Tehran", "Mashhad"]}},
    )  # fmt: skip
    DashboardWidget.objects.create(dashboard=board, saved_report=report)
    DashboardWidget.objects.create(
        dashboard=board,
        title="Orders",
        spec={"semantic_model": "order_metrics", "measures": ["orders"], "chart": "kpi"},
        y=4,
    )
    s = Schedule.objects.create(
        owner=people["owner"], dashboard=board, frequency="weekly", weekday=0,
        time=dt.time(7), timezone="UTC", attachment="csv",
    )  # fmt: skip
    s.recipients.set([people["tehran"], people["east"]])
    scheduling.deliver(s)
    tehran, east = message_for("tehran"), message_for("east")
    assert tehran.subject.startswith("Sales overview · ")
    # the dashboard's default city filter applies, intersected with each scope
    assert csv_cities(east) == {"Mashhad"} or csv_cities(east) == set()
    assert csv_cities(tehran) <= {"Tehran"}
    assert ">Orders<" in html(tehran) and ">Revenue by city<" in html(tehran)


# --- API ------------------------------------------------------------------------------------


def login(user) -> Client:
    client = Client()
    client.force_login(user)
    return client


def test_create_schedule_api(people, report):
    client = login(people["owner"])
    body = {
        "saved_report": report.pk,
        "frequency": "weekly",
        "weekday": 5,
        "time": "08:00",
        "recipients": [people["tehran"].pk, people["east"].pk],
    }
    resp = client.post(SCHEDULES, body, "application/json")
    assert resp.status_code == 201, resp.content
    data = resp.json()
    assert data["target"] == {"kind": "saved_report", "id": report.pk, "name": "Revenue by city"}
    assert data["timezone"] == "Asia/Tehran" and data["calendar"] == "gregorian"
    assert data["attachment"] == "xlsx" and data["next_run_at"]
    assert [r["username"] for r in data["recipients"]] == ["east", "tehran"]  # by name
    nxt = dt.datetime.fromisoformat(data["next_run_at"])
    assert nxt.astimezone(scheduling.zone("Asia/Tehran")).weekday() == 5
    # the owner's time-zone preference is the default
    UserPreferences.objects.create(user=people["owner"], options={"timezone": "Europe/Berlin"})
    resp = client.post(SCHEDULES, body, "application/json")
    assert resp.json()["timezone"] == "Europe/Berlin"


@pytest.mark.parametrize(
    ("change", "field"),
    [
        ({"weekday": None}, "weekday"),
        ({"frequency": "monthly"}, "day_of_month"),
        ({"timezone": "Mars/Olympus"}, "timezone"),
        ({"recipients": []}, "recipients"),
    ],
)
def test_schedule_validation(people, report, change, field):
    body = {
        "saved_report": report.pk,
        "frequency": "weekly",
        "weekday": 0,
        "time": "08:00",
        "recipients": [people["tehran"].pk],
        **change,
    }
    resp = login(people["owner"]).post(SCHEDULES, body, "application/json")
    assert resp.status_code == 400
    assert field in resp.json()["errors"]


def test_recipients_must_be_able_to_open_it(people, report, django_user_model):
    report.visibility = "private"
    report.save()
    client = login(people["owner"])
    body = {"saved_report": report.pk, "frequency": "daily", "time": "08:00"}
    resp = client.post(SCHEDULES, {**body, "recipients": [people["tehran"].pk]}, "application/json")
    assert resp.status_code == 400
    assert "tehran" in resp.json()["errors"]["recipients"][0]
    customer = django_user_model.objects.create_user("customer", email="c@example.com")
    resp = client.post(SCHEDULES, {**body, "recipients": [customer.pk]}, "application/json")
    assert resp.status_code == 400  # not a user of the system (staff)
    # nobody schedules what they can't open
    resp = login(people["tehran"]).post(
        SCHEDULES, {**body, "recipients": [people["tehran"].pk]}, "application/json"
    )
    assert resp.status_code == 404


def test_schedules_are_private_to_their_owner(schedule, people):
    url = f"{SCHEDULES}/{schedule.pk}"
    other = login(people["tehran"])
    assert other.get(url).status_code == 404
    assert other.post(f"{url}/send").status_code == 404
    assert other.get(SCHEDULES).json() == {"schedules": []}
    owner = login(people["owner"])
    resp = owner.patch(url, {"enabled": False, "time": "09:30"}, "application/json")
    assert resp.status_code == 200
    assert resp.json()["enabled"] is False and resp.json()["next_run_at"] is None
    assert resp.json()["time"] == "09:30"
    assert owner.delete(url).json() == {"deleted": True}


def test_send_now(schedule, people):
    owner = login(people["owner"])
    resp = owner.post(f"{SCHEDULES}/{schedule.pk}/send?only_me=true")
    assert [d["user"] for d in resp.json()["deliveries"]] == ["owner"]
    assert [m.to for m in mail.outbox] == [["owner@example.com"]]
    resp = owner.post(f"{SCHEDULES}/{schedule.pk}/send")
    assert len(resp.json()["deliveries"]) == 3 and len(mail.outbox) == 4
    assert ScheduleDelivery.objects.filter(trigger="manual").count() == 4
    detail = owner.get(f"{SCHEDULES}/{schedule.pk}").json()
    assert len(detail["deliveries"]) == 4


def test_delivery_audit_log_needs_the_audit_permission(schedule, people, make_staff):
    scheduling.send_due()
    assert login(people["owner"]).get(f"{API}/audit/deliveries").status_code == 403
    auditor = make_staff("bi_admin.view_reportlog", username="auditor")
    data = login(auditor).get(f"{API}/audit/deliveries", {"status": "sent"}).json()
    assert data["count"] == 3
    assert {e["user"] for e in data["entries"]} == {"owner", "tehran", "east"}


# --- management command ---------------------------------------------------------------------


def test_management_command(schedule):
    out = StringIO()
    call_command("gozar_send_schedules", "--dry-run", stdout=out)
    assert "1 schedule(s) due." in out.getvalue() and not mail.outbox
    out = StringIO()
    call_command("gozar_send_schedules", stdout=out)
    assert "Sent 3 of 3 e-mail(s)." in out.getvalue() and len(mail.outbox) == 3
    out = StringIO()
    call_command("gozar_send_schedules", stdout=out)
    assert "Sent 0 of 0" in out.getvalue()
    call_command("gozar_send_schedules", "--schedule", str(schedule.pk), stdout=StringIO())
    assert len(mail.outbox) == 6


def test_worker_sends_due_schedules(schedule):
    out = StringIO()
    call_command("gozar_worker", "--burst", stdout=out)
    assert len(mail.outbox) == 3 and "E-mail Revenue by city" in out.getvalue()


def test_seeded_demo_schedule(db):
    call_command("seed_demo", customers=10, orders=30, seed=3, users=True, stdout=StringIO())
    s = Schedule.objects.get(name="Weekly sales overview")
    assert s.dashboard.name == "Sales overview" and s.frequency == "weekly"
    assert sorted(u.username for u in s.recipients.all()) == ["east", "manager", "tehran"]
    assert s.next_run_at > timezone.now()
    deliveries = scheduling.deliver(s)
    assert [d.status for d in deliveries] == ["sent"] * 3
    board = Dashboard.objects.get(name="Sales overview")
    assert board.widgets.count() == 5 and board.visibility == "everyone"
    assert SavedReport.objects.filter(owner__username="manager").count() == 4
    # idempotent
    call_command("seed_demo", customers=1, orders=1, seed=3, users=True, stdout=StringIO())
    assert Schedule.objects.count() == 1 and Dashboard.objects.count() == 1
