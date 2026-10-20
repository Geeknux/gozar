"""M6.5: CSV/XLSX export of lists and import with a dry run."""

import csv
import io

import pytest
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile

from bi_admin.calendar import jalali
from bi_admin.models import ChangeLog
from shop.models import City, Customer, Order, Product

API = "/gozar/api"
pytestmark = pytest.mark.django_db
TEHRAN_REGION = ["Tehran", "Karaj", "Qom"]


def crud(label, suffix=""):
    app, model = label.split(".")
    return f"{API}/crud/{app}/{model}/{suffix}"


def read_csv(resp) -> list[list[str]]:
    text = resp.content.decode("utf-8")
    assert text.startswith("﻿")  # BOM
    return list(csv.reader(io.StringIO(text[1:])))


def read_xlsx(resp):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    return wb.active, [list(r) for r in wb.active.iter_rows(values_only=True)]


def csv_file(rows, name="data.csv", encoding="utf-8"):
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return SimpleUploadedFile(name, buf.getvalue().encode(encoding), content_type="text/csv")


def xlsx_file(rows, name="data.xlsx"):
    from openpyxl import Workbook

    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue())


def do_import(client, label, upload, dry_run=True):
    url = crud(label, f"import?dry_run={'true' if dry_run else 'false'}")
    return client.post(url, {"file": upload})


# --- export --------------------------------------------------------------------------


def test_export_csv_follows_filters_ordering_and_labels(admin_client, seeded):
    resp = admin_client.get(
        crud("shop.order", "export"), {"format": "csv", "status": "paid", "ordering": "amount"}
    )
    assert resp.status_code == 200
    assert resp["Content-Disposition"] == 'attachment; filename="order.csv"'
    rows = read_csv(resp)
    assert rows[0] == [
        "Order",
        "Customer",
        "Customer / City / Name",
        "Status",
        "Amount",
        "Paid",
        "Created at",
    ]
    body = rows[1:]
    assert len(body) == Order.objects.filter(status="paid").count()
    assert {r[3] for r in body} == {"Paid"}  # choice labels
    amounts = [int(r[4]) for r in body]
    assert amounts == sorted(amounts)
    assert {r[5] for r in body} == {"True"}


def test_export_xlsx_types_columns_and_selection(admin_client, seeded):
    pks = list(Customer.objects.values_list("pk", flat=True)[:3])
    resp = admin_client.get(
        crud("shop.customer", "export"),
        {"format": "xlsx", "columns": "name,created_at", "pks": ",".join(map(str, pks))},
    )
    assert resp.status_code == 200
    ws, rows = read_xlsx(resp)
    assert rows[0] == ["Name", "Joined at"] and len(rows) == 4
    import datetime as dt

    assert all(isinstance(r[1], dt.datetime) for r in rows[1:])
    assert ws.freeze_panes == "A2" and ws.sheet_view.rightToLeft is False


def test_export_in_persian_is_rtl_with_persian_headers(admin_client, admin_user, seeded):
    from bi_admin.models import UserPreferences

    UserPreferences.objects.create(user=admin_user, language="fa")
    ws, rows = read_xlsx(admin_client.get(crud("shop.city", "export"), {"format": "xlsx"}))
    assert ws.sheet_view.rightToLeft is True
    assert rows[0][1] == "نام"


def test_export_is_scoped_and_hides_restricted_cells(client, regional, seeded):
    client.force_login(regional())
    rows = read_csv(client.get(crud("shop.order", "export"), {"format": "csv"}))
    expected = Order.objects.filter(customer__city__name_en__in=TEHRAN_REGION).count()
    assert len(rows) - 1 == expected
    assert {r[2] for r in rows[1:]} <= set(TEHRAN_REGION)
    # order items list products (unrestricted) but only in-scope orders
    items = read_csv(client.get(crud("shop.orderitem", "export"), {"format": "csv"}))
    from shop.models import OrderItem

    assert (
        len(items) - 1
        == OrderItem.objects.filter(order__customer__city__name_en__in=TEHRAN_REGION).count()
    )


def test_export_restricted_related_cells(client, make_staff, seeded):
    """A customer list cell pointing at a city outside the user's scope says Restricted."""
    from django.contrib.auth.models import Group

    user = make_staff("shop.view_customer", "shop.view_city", username="mixed")
    user.groups.add(Group.objects.get_or_create(name="Sales: Tehran region")[0])
    client.force_login(user)
    # the customer list itself is limited by region too; drop that rule for this check
    from shop.gozar import CustomerOptions

    rules = CustomerOptions.row_rules
    CustomerOptions.row_rules = []
    try:
        rows = read_csv(client.get(crud("shop.customer", "export"), {"format": "csv"}))
    finally:
        CustomerOptions.row_rules = rules
    cities = {r[2] for r in rows[1:]}
    assert "Restricted" in cities and cities - {"Restricted"} <= set(TEHRAN_REGION)


def test_export_permissions_formats_and_truncation(make_staff, admin_client, seeded, settings):
    from django.test import Client

    client = Client()
    client.force_login(make_staff("shop.view_city"))
    assert client.get(crud("shop.order", "export")).status_code == 403
    assert client.get(crud("shop.city", "export"), {"format": "pdf"}).status_code == 400
    settings.GOZAR = {**settings.GOZAR, "CRUD_EXPORT_MAX_ROWS": 5}
    resp = admin_client.get(crud("shop.order", "export"), {"format": "csv"})
    assert resp["X-Gozar-Truncated"] == "true" and len(read_csv(resp)) == 6


def test_import_template(admin_client, seeded):
    rows = read_csv(admin_client.get(crud("shop.city", "import-template"), {"format": "csv"}))
    assert rows == [["id", "name", "name_en", "province"]]
    ws, rows = read_xlsx(
        admin_client.get(crud("shop.product", "import-template"), {"format": "xlsx"})
    )
    assert "datasheet" not in rows[0] and "tags" in rows[0]


# --- import --------------------------------------------------------------------------


def test_import_dry_run_reports_and_saves_nothing(admin_client, seeded):
    tehran = City.objects.get(name_en="Tehran")
    before = Customer.objects.count()
    upload = csv_file(
        [
            ["Name", "Phone", "City", "Joined at"],
            ["New Person", "0912", "Tehran", "2026-01-02 10:00"],
            ["Other Person", "0913", str(tehran.pk), "1404/10/12"],  # pk + Jalali date
        ]
    )
    resp = do_import(admin_client, "shop.customer", upload)
    assert resp.status_code == 200, resp.json()
    report = resp.json()
    assert report["dry_run"] is True and report["committed"] is False
    assert report["totals"] == {"create": 2, "update": 0, "skip": 0, "error": 0}
    assert [c["field"] for c in report["columns"]] == ["name", "phone", "city", "created_at"]
    assert report["rows"][0] == {"row": 2, "action": "create", "pk": None, "label": "New Person"}
    assert Customer.objects.count() == before
    assert not ChangeLog.objects.filter(changes__source="import").exists()


def test_import_commits_creates_and_updates_with_history(admin_client, seeded):
    customer = Customer.objects.first()
    upload = xlsx_file(
        [
            ["ID", "name", "phone", "city", "created_at"],
            [
                customer.pk,
                "Renamed",
                customer.phone,
                customer.city.name,
                customer.created_at.replace(tzinfo=None),
            ],
            [None, "Brand New", "0999", "Shiraz", "2026-05-05"],
        ]
    )
    resp = do_import(admin_client, "shop.customer", upload, dry_run=False)
    assert resp.status_code == 200, resp.json()
    report = resp.json()
    assert report["committed"] is True
    assert report["totals"]["create"] == 1 and report["totals"]["update"] == 1
    customer.refresh_from_db()
    assert customer.name == "Renamed"
    new = Customer.objects.get(name="Brand New")
    assert new.city.name_en == "Shiraz"
    assert report["rows"][1]["pk"] == str(new.pk)
    entry = ChangeLog.objects.get(model_label="shop.customer", object_pk=str(customer.pk))
    assert entry.changes["source"] == "import" and "name" in entry.changes["fields"]


def test_import_unchanged_rows_are_skipped(admin_client, seeded):
    city = City.objects.first()
    upload = csv_file(
        [["id", "name", "name_en", "province"], [city.pk, city.name, city.name_en, city.province]]
    )
    report = do_import(admin_client, "shop.city", upload, dry_run=False).json()
    assert report["totals"] == {"create": 0, "update": 0, "skip": 1, "error": 0}


def test_import_errors_roll_everything_back(admin_client, seeded):
    before = Product.objects.count()
    upload = csv_file(
        [
            ["name", "category", "price", "is_active", "tags", "unknown column"],
            ["Good", "Books", "100", "yes", "", "x"],
            ["Bad price", "books", "abc", "no", "", "x"],
            ["Bad category", "toys", "1", "1", "", "x"],
            ["Bad bool", "grocery", "1", "maybe", "", "x"],
        ]
    )
    report = do_import(admin_client, "shop.product", upload, dry_run=False).json()
    assert report["committed"] is False
    assert report["totals"] == {"create": 1, "update": 0, "skip": 0, "error": 3}
    assert report["ignored_columns"] == ["unknown column"]
    errors = {r["row"]: r["errors"] for r in report["rows"] if r["action"] == "error"}
    assert "price" in errors[3] and "category" in errors[4] and "is_active" in errors[5]
    assert Product.objects.count() == before  # the good row was rolled back too


def test_import_choice_labels_persian_headers_and_digits(admin_client, seeded):
    upload = csv_file(
        [
            ["نام", "دسته", "قیمت", "فعال"],
            ["کتاب نو", "کتاب", "۱۲۰٬۰۰۰", "بله"],
        ]
    )
    report = do_import(admin_client, "shop.product", upload, dry_run=False).json()
    assert report["committed"], report
    p = Product.objects.get(name="کتاب نو")
    assert (p.category, p.price, p.is_active) == ("books", 120000, True)


def test_import_relation_labels_must_be_unambiguous_and_in_scope(client, regional, seeded):
    user = regional()
    user.user_permissions.add(*Permission.objects.filter(codename__in=["add_customer"]))
    client.force_login(user)
    upload = csv_file(
        [
            ["name", "phone", "city", "created_at"],
            ["In region", "1", "Karaj", "2026-01-01"],
            ["Outside", "2", "Shiraz", "2026-01-01"],
            ["Missing", "3", "Atlantis", "2026-01-01"],
        ]
    )
    report = do_import(client, "shop.customer", upload).json()
    assert [r["action"] for r in report["rows"]] == ["create", "error", "error"]
    # a city outside the scope is "not found", exactly like one that doesn't exist
    assert report["rows"][1]["errors"]["city"] == ["“Shiraz” was not found."]
    assert report["rows"][2]["errors"]["city"] == ["“Atlantis” was not found."]


def test_import_updates_only_rows_in_scope(client, regional, seeded):
    user = regional()
    user.user_permissions.add(*Permission.objects.filter(codename__in=["change_customer"]))
    client.force_login(user)
    outside = Customer.objects.exclude(city__name_en__in=TEHRAN_REGION).first()
    upload = csv_file([["id", "name"], [outside.pk, "Hijacked"]])
    report = do_import(client, "shop.customer", upload, dry_run=False).json()
    assert report["totals"]["error"] == 1 and not report["committed"]
    assert "ID" in report["rows"][0]["errors"]["__all__"][0]
    outside.refresh_from_db()
    assert outside.name != "Hijacked"


def test_import_permissions_and_bad_files(make_staff, admin_client, seeded, settings):
    from django.test import Client

    client = Client()
    client.force_login(make_staff("shop.view_city"))
    upload = csv_file([["name", "province"], ["X", "Y"]])
    assert do_import(client, "shop.city", upload).status_code == 403

    client.force_login(make_staff("shop.change_city", username="changer"))
    report = do_import(client, "shop.city", csv_file([["name", "province"], ["X", "Y"]])).json()
    assert report["rows"][0]["errors"]["__all__"] == ["You may not add records."]

    assert admin_client.post(crud("shop.city", "import")).status_code == 400
    empty = do_import(admin_client, "shop.city", csv_file([]))
    assert empty.status_code == 400
    nothing = do_import(admin_client, "shop.city", csv_file([["foo"], ["bar"]]))
    assert nothing.status_code == 400
    broken = do_import(admin_client, "shop.city", SimpleUploadedFile("x.xlsx", b"PK\x03\x04junk"))
    assert broken.status_code == 400
    settings.GOZAR = {**settings.GOZAR, "CRUD_IMPORT_MAX_ROWS": 1}
    big = do_import(
        admin_client, "shop.city", csv_file([["name", "province"], ["A", "B"], ["C", "D"]])
    )
    assert big.status_code == 400 and big.json()["code"] == "too_many_rows"


def test_import_semicolon_csv_and_m2m(admin_client, seeded):
    from shop.models import Tag

    a, b = Tag.objects.create(name="sale"), Tag.objects.create(name="new")
    text = "name;category;price;tags\nTagged;home;5;sale, new\n"
    upload = SimpleUploadedFile("t.csv", text.encode("utf-8"))
    report = do_import(admin_client, "shop.product", upload, dry_run=False).json()
    assert report["committed"], report
    assert set(Product.objects.get(name="Tagged").tags.all()) == {a, b}


def test_jalali_dates_import_to_gregorian(admin_client, seeded):
    upload = csv_file(
        [["name", "phone", "city", "created_at"], ["J", "1", "Tehran", "۱۴۰۵/۰۷/۲۶ ۰۸:۳۰"]]
    )
    report = do_import(admin_client, "shop.customer", upload, dry_run=False).json()
    assert report["committed"], report
    created = Customer.objects.get(name="J").created_at
    from django.utils import timezone

    local = timezone.localtime(created)
    assert local.date() == jalali.to_gregorian(1405, 7, 26) and (local.hour, local.minute) == (
        8,
        30,
    )
